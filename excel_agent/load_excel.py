# import pandas as pd



# df=pd.read_excel(r"data/Question1/Inventory-Records-Sample-Data.xlsx")


import openpyxl

# Load the workbook and select the active sheet
workbook = openpyxl.load_workbook('data/Question1/Inventory-Records-Sample-Data.xlsx')
sheet = workbook.active

# 1. Sheets
sheets = workbook.sheetnames

# 2 & 3. Rows and Columns
total_rows = sheet.max_row
total_cols = sheet.max_column

# Find the first row with more than one populated cell, skipping title rows
header_row = next(
    (
        row_idx
        for row_idx in range(1, total_rows + 1)
        if sum(sheet.cell(row=row_idx, column=col_idx).value is not None
               for col_idx in range(1, total_cols + 1)) > 1
    ),
    None,
)
if header_row is None:
    header_row = next(
        row_idx
        for row_idx in range(1, total_rows + 1)
        if any(sheet.cell(row=row_idx, column=col_idx).value is not None
               for col_idx in range(1, total_cols + 1))
    )

# Get column names from the header row, keeping their Excel positions
columns = [
    (col_idx, sheet.cell(row=header_row, column=col_idx).value)
    for col_idx in range(1, total_cols + 1)
    if sheet.cell(row=header_row, column=col_idx).value is not None
]

# 5 & 6. Analyze Data Types and Missing Values (Checking the first data row)
sample_types = {}
missing_counts = {col_name: 0 for _, col_name in columns}

for col_idx, col_name in columns:
    # Get a sample data type from the first data row
    sample_val = sheet.cell(row=header_row + 1, column=col_idx).value
    sample_types[col_name] = type(sample_val).__name__ if sample_val is not None else "Unknown"
    
    # Count missing values in this column
    for row_idx in range(header_row + 1, total_rows + 1):
        if sheet.cell(row=row_idx, column=col_idx).value is None:
            missing_counts[col_name] += 1

# Print the Results
print(f"Sheets: {', '.join(sheets)}")
print(f"Rows: {total_rows - header_row} (excluding header)")
print(f"Columns: {len(columns)} ({', '.join(str(name) for _, name in columns)})")
print("\nData types (Sample from first data row):")
for col, t in sample_types.items():
    print(f"  - {col}: {t}")
print("\nMissing values:")
for col, count in missing_counts.items():
    print(f"  - {col}: {count} empty cells")

product_name_col = next(
    (
        col_idx
        for col_idx, col_name in columns
        if "".join(char.lower() for char in str(col_name) if char.isalnum()) == "productname"
    ),
    None,
)
inventory_col = next(
    (
        col_idx
        for col_idx, col_name in columns
        if "".join(char.lower() for char in str(col_name) if char.isalnum()) == "handinstock"
    ),
    None,
)

if product_name_col is None or inventory_col is None:
    raise ValueError("Cannot find Product Name and Hand-In-Stock columns.")

inventories = [
    (sheet.cell(row=row_idx, column=product_name_col).value,
     sheet.cell(row=row_idx, column=inventory_col).value)
    for row_idx in range(header_row + 1, total_rows + 1)
    if sheet.cell(row=row_idx, column=product_name_col).value is not None
    and sheet.cell(row=row_idx, column=inventory_col).value is not None
]
if not inventories:
    raise ValueError("No product inventory values were found.")

highest_inventory = max(quantity for _, quantity in inventories)
top_products = [
    name for name, quantity in inventories if quantity == highest_inventory
]
print(
    "\nProduct with highest inventory (by Hand-In-Stock): "
    f"{', '.join(str(name) for name in top_products)} "
    f"({highest_inventory} units)"
)